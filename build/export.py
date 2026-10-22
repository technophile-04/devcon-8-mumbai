"""Exports a kit deck for Devcon's Google Slides.

Run: build/export.sh decks/<talk-slug>.html [outdir]
     build/export.sh path/to/slide-images/ [outdir]

From an HTML deck (a copy of reference/deck.html) it writes, next to the deck unless outdir is given:
  <slug>-slides/NN-name.png   one 1920x1080 render per slide, by Google Chrome, like build/render.sh
  <slug>-google-slides.pptx   one full-bleed picture per slide, 10 x 5.625 in (Devcon's deck size), speaker notes
  <slug>.pdf                  the same images, the laptop backup
Slides are counted from the deck: every <section class="slide">, in document order. Speaker notes are the
text of an <aside class="notes"> inside the section. The export stops if PP Neue Machina, Space Grotesk or
Space Mono did not load, because the pictures would carry the wrong font.

From a folder of images (a deck built by hand in Keynote from the pptx template, then exported as images)
it takes the images in name order and writes only the pptx and the PDF, without notes.
"""
import os
import re
import subprocess
import sys
from html.parser import HTMLParser
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

FONTS = ("PP Neue Machina", "Space Grotesk", "Space Mono")
W, H = 1920, 1080
PAGE_W, PAGE_H = 10, 5.625          # in; Devcon's Google Slides decks are 16:9 at this size
CHROMES = (
    os.environ.get("CHROME"),
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "google-chrome", "chromium", "chromium-browser",
)


def die(msg):
    sys.exit(f"export: {msg}")


def chrome():
    from shutil import which
    for c in CHROMES:
        if c and (Path(c).is_file() or which(c)):
            return c
    die("Google Chrome not found; set CHROME=/path/to/chrome")


def run_chrome(*args):
    out = subprocess.run(
        [chrome(), "--headless=new", "--disable-gpu", "--hide-scrollbars",
         "--force-device-scale-factor=1", "--virtual-time-budget=10000", *args],
        capture_output=True, text=True)
    if out.returncode:
        die(f"chrome failed: {out.stderr.strip()[-400:]}")
    return out.stdout


class Deck(HTMLParser):
    """Reads the rendered DOM: loaded fonts from <body data-fonts>, slides from <section class="slide">,
    notes from <aside class="notes"> inside a slide."""

    def __init__(self):
        super().__init__()
        self.fonts, self.slides, self._in_notes = [], [], False

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        classes = (a.get("class") or "").split()
        if tag == "body":
            self.fonts = [f.strip() for f in (a.get("data-fonts") or "").split(",") if f.strip()]
        elif tag == "section" and "slide" in classes:
            self.slides.append({"name": a.get("data-name") or "slide", "notes": []})
        elif tag == "aside" and "notes" in classes and self.slides:
            self._in_notes = True

    def handle_endtag(self, tag):
        if tag == "aside":
            self._in_notes = False

    def handle_data(self, data):
        if self._in_notes:
            self.slides[-1]["notes"].append(data)


def notes_text(parts):
    text = "".join(parts)
    return "\n".join(line for line in (re.sub(r"[ \t]+", " ", l).strip() for l in text.splitlines()) if line)


def render_html(deck, outdir):
    url = deck.resolve().as_uri()
    parsed = Deck()
    parsed.feed(run_chrome("--dump-dom", url))
    missing = [f for f in FONTS if f not in parsed.fonts]
    if missing:
        die(f"{', '.join(missing)} did not load (loaded: {', '.join(parsed.fonts) or 'none'}). "
            "Check fonts/ and the Google Fonts link in the deck, and that you are online.")
    if not parsed.slides:
        die("no <section class=\"slide\"> in the deck")
    slides_dir = outdir / f"{deck.stem}-slides"
    slides_dir.mkdir(parents=True, exist_ok=True)
    for old in slides_dir.glob("*.png"):
        old.unlink()
    pngs, notes = [], []
    for i, s in enumerate(parsed.slides, 1):
        png = slides_dir / f"{i:02d}-{s['name']}.png"
        run_chrome(f"--window-size={W},{H}", f"--screenshot={png}", f"{url}?slide={i}")
        with Image.open(png) as im:
            if im.size != (W, H):
                die(f"{png.name} is {im.size[0]}x{im.size[1]}, not {W}x{H}")
        pngs.append(png)
        notes.append(notes_text(s["notes"]))
        print(f"rendered {png.relative_to(outdir)}" + ("  (notes)" if notes[-1] else ""))
    return pngs, notes


def read_folder(folder):
    pngs = sorted(p for p in folder.iterdir() if p.suffix.lower() in (".png", ".jpg", ".jpeg"))
    if not pngs:
        die(f"no png or jpg images in {folder}")
    for p in pngs:
        with Image.open(p) as im:
            w, h = im.size
        if abs(w / h - 16 / 9) > 0.01:
            die(f"{p.name} is {w}x{h}, not 16:9; export slides at 1920x1080")
    return pngs, [""] * len(pngs)


def write_pptx(pngs, notes, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(PAGE_W), Inches(PAGE_H)
    blank = prs.slide_layouts[6]
    for png, note in zip(pngs, notes):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
        if note:
            slide.notes_slide.notes_text_frame.text = note
    prs.save(path)


def write_pdf(pngs, path):
    pages = [Image.open(p).convert("RGB") for p in pngs]
    pages[0].save(path, save_all=True, append_images=pages[1:], resolution=W / PAGE_W)


def main(argv):
    if len(argv) not in (2, 3):
        die("usage: build/export.sh <deck.html | folder-of-images> [outdir]")
    src = Path(argv[1])
    if not src.exists():
        die(f"{src} does not exist")
    outdir = Path(argv[2]) if len(argv) == 3 else src.parent
    outdir.mkdir(parents=True, exist_ok=True)
    if src.is_dir():
        pngs, notes = read_folder(src)
    else:
        pngs, notes = render_html(src, outdir)
    pptx = outdir / f"{src.stem}-google-slides.pptx"
    pdf = outdir / f"{src.stem}.pdf"
    write_pptx(pngs, notes, pptx)
    write_pdf(pngs, pdf)
    mb = lambda p: f"{p.stat().st_size / 1e6:.1f} MB"
    print(f"{len(pngs)} slides, {sum(1 for n in notes if n)} with notes")
    print(f"  {pptx}  {mb(pptx)}  import this into your Devcon Google Slides deck")
    print(f"  {pdf}  {mb(pdf)}  laptop backup")


if __name__ == "__main__":
    main(sys.argv)
